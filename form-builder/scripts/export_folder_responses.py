import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def main():
    input_path, output_path = sys.argv[1:3]
    payload = json.loads(open(input_path, "r", encoding="utf-8").read())
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Folder responses"

    columns = payload.get("columns", [])
    rows = payload.get("rows", [])
    headers = [item.get("label") or item.get("key") or "Column" for item in columns]
    keys = [item.get("key") or "" for item in columns]
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(key, "") for key in keys])

    header_fill = PatternFill("solid", fgColor="173B68")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32

    for index, column in enumerate(sheet.columns, 1):
        longest = max((len(str(cell.value or "")) for cell in column), default=0)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 12), 42)
        for cell in column[1:]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    workbook.save(output_path)


if __name__ == "__main__":
    main()
