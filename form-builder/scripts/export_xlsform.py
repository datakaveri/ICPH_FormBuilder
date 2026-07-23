#!/usr/bin/env python3
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


SURVEY_HEADERS = [
    "type",
    "name",
    "label",
    "hint",
    "required",
    "relevant",
    "appearance",
    "default",
    "constraint",
    "constraint_message",
    "calculation",
    "trigger",
    "choice_filter",
    "parameters",
    "repeat_count",
    "note",
    "image",
    "audio ",
    "video",
    "read_only",
    "required_message",
    "guidance_hint",
    "save_to",
    "big-image",
]

CHOICES_HEADERS = ["list_name", "name", "label", "image", "audio", "video", "big-image", "geometry"]
SETTINGS_HEADERS = [
    "form_title",
    "form_id",
    "version",
    "instance_name",
    "default_language",
    "style",
    "public_key",
    "submission_url",
    "allow_choice_duplicates",
]
ENTITIES_HEADERS = ["list_name", "label", "create_if", "entity_id", "update_if"]


def slug(value: str) -> str:
    text = re.sub(r"[^A-Za-z0-9_]+", "_", str(value or "")).strip("_")
    if not text:
        text = "field"
    if text[0].isdigit():
        text = f"q_{text}"
    return text[:64]


def odk_name(value: str, fallback: str = "field") -> str:
    text = str(value or "").strip()
    if re.match(r"^[A-Za-z_][A-Za-z0-9_.-]*$", text):
        return text[:64]
    return slug(text or fallback)


def list_name_for(question: dict) -> str:
    fallback = f"{question.get('name', 'choice')}_choices"
    return odk_name(question.get("listName") or fallback, fallback)


def choice_name_for(option: dict, index: int) -> str:
    name = str(option.get("name") or option.get("value") or "").strip()
    if name:
        return name
    label = str(option.get("label") or "").strip()
    return slug(label or f"choice_{index + 1}")


def choice_signature(row: dict) -> tuple:
    extras = tuple(sorted((row.get("extraColumns") or {}).items()))
    return (
        row.get("label") or "",
        row.get("image") or "",
        row.get("audio") or "",
        row.get("video") or "",
        row.get("big-image") or "",
        row.get("geometry") or "",
        extras,
    )


def xls_type(question: dict) -> str:
    qtype = question.get("type", "text")
    if qtype in {"select_one", "select_multiple", "rank"}:
        return f"{qtype} {list_name_for(question)}"
    if qtype in {"select_one_from_file", "select_multiple_from_file"}:
        return f"{qtype} {question.get('listName') or 'choices.csv'}"
    if qtype == "dateTime":
        return "datetime"
    return qtype


def write_header(ws, headers):
    fill = PatternFill("solid", fgColor="EAF2F8")
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = fill
        ws.column_dimensions[get_column_letter(idx)].width = max(14, min(36, len(header) + 6))


def ordered_headers(base_headers, rows, extra_keys=("extraColumns",)):
    headers = list(base_headers)
    seen = set(headers)
    for row in rows:
        for key in extra_keys:
            for header in (row.get(key) or {}).keys():
                if header and header not in seen:
                    headers.append(header)
                    seen.add(header)
    return headers


def valid_instance_name(value: str, survey_names: set[str]) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    refs = re.findall(r"\$\{([^}]+)\}", text)
    if refs and any(ref not in survey_names for ref in refs):
        return None
    return text


def boolean_or_expression(flag, expression=None, force_yes=False):
    if force_yes:
        return "yes"
    text = str(expression or "").strip()
    if text:
        return text
    return "yes" if flag else None


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: export_xlsform.py draft.json output.xlsx", file=sys.stderr)
        return 2

    draft_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    form = json.loads(draft_path.read_text(encoding="utf-8"))
    questions = list(form.get("questions") or [])
    survey_names = {odk_name(question.get("name") or question.get("id")) for question in questions}

    wb = Workbook()
    survey = wb.active
    survey.title = "survey"
    choices = wb.create_sheet("choices")
    settings = wb.create_sheet("settings")
    entities = wb.create_sheet("entities")

    survey_headers = ordered_headers(SURVEY_HEADERS, questions)
    choice_rows = []
    choices_by_list = {}
    for question in questions:
        if question.get("type") not in {"select_one", "select_multiple", "rank"}:
            continue
        list_name = list_name_for(question)
        list_choices = choices_by_list.setdefault(list_name, {})
        for option_index, option in enumerate(question.get("options") or []):
            name = choice_name_for(option, option_index)
            label = str(option.get("label") or "").strip()
            if not name or not label:
                continue
            row = {
                "list_name": list_name,
                "name": name,
                "label": label,
                "image": option.get("image") or None,
                "audio": option.get("audio") or None,
                "video": option.get("video") or None,
                "big-image": option.get("bigImage") or option.get("big-image") or None,
                "geometry": option.get("geometry") or None,
                "extraColumns": option.get("extraColumns") or {},
            }
            existing = list_choices.get(name)
            if existing:
                # Shared XLSForm choice lists must be written once. Imported
                # forms often reuse the same list on many questions.
                if choice_signature(existing) != choice_signature(row):
                    print(
                        f"Warning: duplicate choice '{name}' in list '{list_name}' has conflicting values; keeping the first.",
                        file=sys.stderr,
                    )
                continue
            list_choices[name] = row
            choice_rows.append(row)
    choice_headers = ordered_headers(CHOICES_HEADERS, choice_rows)
    settings_extra = form.get("settingsExtraColumns") or {}
    settings_headers = list(SETTINGS_HEADERS)
    for header in settings_extra.keys():
        if header not in settings_headers:
            settings_headers.append(header)
    entity_rows = form.get("entities") if isinstance(form.get("entities"), list) else []
    entity_headers = ordered_headers(ENTITIES_HEADERS, entity_rows)

    write_header(survey, survey_headers)
    write_header(choices, choice_headers)
    write_header(settings, settings_headers)
    write_header(entities, entity_headers)

    for row_idx, question in enumerate(questions, start=2):
        values = {
            **(question.get("extraColumns") or {}),
            "type": xls_type(question),
            "name": odk_name(question.get("name") or question.get("id")),
            "label": question.get("label") or question.get("name") or "Untitled question",
            "hint": question.get("hint") or None,
            "required": boolean_or_expression(question.get("required"), question.get("requiredExpression") or question.get("required_expression")),
            "relevant": question.get("relevant") or None,
            "appearance": question.get("appearance") or None,
            "default": question.get("defaultValue") or None,
            "constraint": question.get("constraint") or None,
            "constraint_message": question.get("constraintMessage") or None,
            "calculation": question.get("calculation") or None,
            "trigger": question.get("trigger") or None,
            "choice_filter": question.get("choiceFilter") or question.get("choice_filter") or None,
            "parameters": question.get("parameters") or None,
            "repeat_count": question.get("repeatCount") or question.get("repeat_count") or None,
            "note": question.get("note") or None,
            "image": question.get("image") or None,
            "audio ": question.get("audio") or question.get("audio ") or None,
            "video": question.get("video") or None,
            "read_only": boolean_or_expression(
                question.get("readOnly"),
                question.get("readOnlyExpression") or question.get("read_only_expression"),
                force_yes=question.get("type") == "calculate",
            ),
            "required_message": question.get("requiredMessage") or question.get("required_message") or None,
            "guidance_hint": question.get("guidanceHint") or question.get("guidance_hint") or None,
            "save_to": question.get("saveTo") or question.get("save_to") or None,
            "big-image": question.get("bigImage") or question.get("big-image") or None,
        }
        for col_idx, header in enumerate(survey_headers, start=1):
            survey.cell(row=row_idx, column=col_idx, value=values.get(header))

    choice_row = 2
    for row in choice_rows:
        values = {**(row.get("extraColumns") or {}), **row}
        for col_idx, header in enumerate(choice_headers, start=1):
            choices.cell(row=choice_row, column=col_idx, value=values.get(header))
        choice_row += 1

    settings_values = {
        **settings_extra,
        "form_title": form.get("title") or "Untitled ICPH Form",
        "form_id": odk_name(form.get("formId") or form.get("title") or "icph_form", "icph_form"),
        "version": form.get("version") or "1",
        "instance_name": valid_instance_name(form.get("instanceName"), survey_names),
        "default_language": form.get("defaultLanguage") or "english",
        "style": form.get("style") or None,
        "public_key": form.get("publicKey") or form.get("public_key") or None,
        "submission_url": form.get("submissionUrl") or form.get("submission_url") or None,
        "allow_choice_duplicates": form.get("allowChoiceDuplicates") or form.get("allow_choice_duplicates") or None,
    }
    for col_idx, header in enumerate(settings_headers, start=1):
        settings.cell(row=2, column=col_idx, value=settings_values.get(header))

    for row_idx, row in enumerate(entity_rows, start=2):
        values = {**(row.get("extraColumns") or {}), **row}
        for col_idx, header in enumerate(entity_headers, start=1):
            entities.cell(row=row_idx, column=col_idx, value=values.get(header))

    for ws in (survey, choices, settings, entities):
        ws.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"Wrote XLSForm: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
