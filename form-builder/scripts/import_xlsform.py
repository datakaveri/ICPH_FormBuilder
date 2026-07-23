#!/usr/bin/env python3
from __future__ import annotations

import json
import re
import sys
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

HEADER_ALIASES = {
    "list name": "list_name",
    "form title": "form_title",
    "title": "form_title",
    "form id": "form_id",
    "id string": "form_id",
    "id_string": "form_id",
    "default language": "default_language",
    "choice filter": "choice_filter",
    "constraint message": "constraint_message",
    "required message": "required_message",
    "guidance hint": "guidance_hint",
    "read only": "read_only",
    "repeat count": "repeat_count",
    "save to": "save_to",
}


def slug(value: str) -> str:
    text = re.sub(r"[^A-Za-z0-9_]+", "_", str(value or "")).strip("_")
    if not text:
        text = "field"
    if text[0].isdigit():
        text = f"q_{text}"
    return text[:64]


def canonical_header(value) -> str:
    header = str(value or "").strip()
    key = " ".join(header.lower().replace("_", " ").split())
    return HEADER_ALIASES.get(key, header)


def rows_as_dicts(ws):
    headers = [canonical_header(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for excel_row in ws.iter_rows(min_row=2, values_only=True):
        if not any(value is not None and str(value).strip() for value in excel_row):
            continue
        row = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            row[header] = excel_row[index] if index < len(excel_row) else None
        rows.append(row)
    return rows


SURVEY_STANDARD_COLUMNS = {
    "type",
    "name",
    "label",
    "hint",
    "required",
    "relevant",
    "appearance",
    "default",
    "constraint",
    "constraint_message",
    "calculation",
    "trigger",
    "choice_filter",
    "parameters",
    "repeat_count",
    "note",
    "image",
    "audio",
    "video",
    "read_only",
    "required_message",
    "guidance_hint",
    "save_to",
    "big-image",
}

CHOICE_STANDARD_COLUMNS = {"list_name", "name", "label", "image", "audio", "video", "big-image", "geometry"}
SETTINGS_STANDARD_COLUMNS = {
    "form_title",
    "form_id",
    "version",
    "instance_name",
    "default_language",
    "style",
    "public_key",
    "submission_url",
    "allow_choice_duplicates",
}
ENTITY_STANDARD_COLUMNS = {"list_name", "label", "create_if", "entity_id", "update_if"}


def clean_value(value) -> str:
    return "" if value is None else str(value).strip()


def json_value(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def extra_columns(row, standard_columns):
    return {
        header: json_value(value)
        for header, value in row.items()
        if header not in standard_columns and value is not None and str(value).strip() != ""
    }


def parse_type(raw_type):
    value = str(raw_type or "text").strip()
    normalized = " ".join(value.lower().replace("_", " ").split())
    match = re.match(r"^(select_one_from_file|select_multiple_from_file|select_one|select_multiple|rank)\s+(.+)$", value, re.I)
    if match:
        return match.group(1).lower(), match.group(2).strip()
    if normalized == "begin group":
        return "begin_group", ""
    if normalized == "end group":
        return "end_group", ""
    if normalized == "begin repeat":
        return "begin_repeat", ""
    if normalized == "end repeat":
        return "end_repeat", ""
    if value == "datetime":
        return "dateTime", ""
    return value, ""


def truthy(value) -> bool:
    return str(value or "").strip().lower() in {"yes", "true", "1"}


def boolean_expression(value) -> str:
    text = clean_value(value)
    if not text or text.lower() in {"yes", "true", "1", "no", "false", "0"}:
        return ""
    return text


def first_column_value(row, base_name):
    direct = clean_value(row.get(base_name))
    if direct:
        return direct
    prefix = f"{base_name}::"
    for header, value in row.items():
        if str(header).startswith(prefix):
            translated = clean_value(value)
            if translated:
                return translated
    return ""


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: import_xlsform.py input.xlsx output-draft.json", file=sys.stderr)
        return 2

    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    wb = load_workbook(input_path, data_only=False)
    if "survey" not in wb.sheetnames:
        raise SystemExit("Missing required XLSForm sheet: survey")

    survey_rows = rows_as_dicts(wb["survey"])
    choice_rows = rows_as_dicts(wb["choices"]) if "choices" in wb.sheetnames else []
    settings_rows = rows_as_dicts(wb["settings"]) if "settings" in wb.sheetnames else []
    entity_rows = rows_as_dicts(wb["entities"]) if "entities" in wb.sheetnames else []

    choices_by_list = {}
    for row in choice_rows:
        list_name = str(row.get("list_name") or "").strip()
        if not list_name:
            continue
        choices_by_list.setdefault(list_name, []).append(
            {
                "id": f"{list_name}_{len(choices_by_list.get(list_name, [])) + 1}",
                "name": clean_value(row.get("name")),
                "label": first_column_value(row, "label"),
                "image": first_column_value(row, "image"),
                "audio": first_column_value(row, "audio"),
                "video": first_column_value(row, "video"),
                "bigImage": clean_value(row.get("big-image")),
                "geometry": clean_value(row.get("geometry")),
                "extraColumns": extra_columns(row, CHOICE_STANDARD_COLUMNS),
            }
        )

    questions = []
    for index, row in enumerate(survey_rows, start=1):
        raw_type, list_name = parse_type(row.get("type"))
        name = str(row.get("name") or f"q{index}").strip()
        question = {
            "id": f"imported_{index}_{slug(name)}",
            "type": raw_type,
            "name": name,
            "label": first_column_value(row, "label") or name,
            "hint": first_column_value(row, "hint"),
            "required": truthy(row.get("required")),
            "requiredExpression": boolean_expression(row.get("required")),
            "relevant": clean_value(row.get("relevant")),
            "appearance": clean_value(row.get("appearance")),
            "defaultValue": clean_value(row.get("default")),
            "constraint": clean_value(row.get("constraint")),
            "constraintMessage": first_column_value(row, "constraint_message"),
            "calculation": clean_value(row.get("calculation")),
            "trigger": clean_value(row.get("trigger")),
            "choiceFilter": clean_value(row.get("choice_filter")),
            "parameters": clean_value(row.get("parameters")),
            "repeatCount": clean_value(row.get("repeat_count")),
            "note": clean_value(row.get("note")),
            "image": first_column_value(row, "image"),
            "audio": first_column_value(row, "audio"),
            "video": first_column_value(row, "video"),
            "requiredMessage": first_column_value(row, "required_message"),
            "guidanceHint": first_column_value(row, "guidance_hint"),
            "saveTo": clean_value(row.get("save_to")),
            "bigImage": clean_value(row.get("big-image")),
            "readOnly": truthy(row.get("read_only")),
            "readOnlyExpression": boolean_expression(row.get("read_only")),
            "extraColumns": extra_columns(row, SURVEY_STANDARD_COLUMNS),
        }
        if raw_type in {"select_one", "select_multiple", "rank"}:
            question["listName"] = list_name
            question["options"] = choices_by_list.get(list_name, [])
        if raw_type in {"select_one_from_file", "select_multiple_from_file"}:
            question["listName"] = list_name
        questions.append(question)

    settings = settings_rows[0] if settings_rows else {}
    title = str(settings.get("form_title") or input_path.stem).strip()
    entities = []
    for row in entity_rows:
        entities.append({
            "list_name": clean_value(row.get("list_name")),
            "label": clean_value(row.get("label")),
            "create_if": clean_value(row.get("create_if")),
            "entity_id": clean_value(row.get("entity_id")),
            "update_if": clean_value(row.get("update_if")),
            "extraColumns": extra_columns(row, ENTITY_STANDARD_COLUMNS),
        })
    form = {
        "title": title,
        "formId": slug(str(settings.get("form_id") or input_path.stem).strip()).lower(),
        "version": str(settings.get("version") or "1").strip(),
        "instanceName": str(settings.get("instance_name") or "").strip(),
        "defaultLanguage": str(settings.get("default_language") or "english").strip(),
        "style": clean_value(settings.get("style")),
        "publicKey": clean_value(settings.get("public_key")),
        "submissionUrl": clean_value(settings.get("submission_url")),
        "allowChoiceDuplicates": clean_value(settings.get("allow_choice_duplicates")),
        "settingsExtraColumns": extra_columns(settings, SETTINGS_STANDARD_COLUMNS),
        "entities": entities,
        "questions": questions,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(form, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Imported XLSForm draft: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
