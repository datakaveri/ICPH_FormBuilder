#!/usr/bin/env python3
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

HEADER_ALIASES = {
    "form title": "form_title",
    "title": "form_title",
    "form id": "form_id",
    "id string": "form_id",
    "id_string": "form_id",
    "list name": "list_name",
}


def canonical_header(value) -> str:
    header = str(value or "").strip()
    key = " ".join(header.lower().replace("_", " ").split())
    return HEADER_ALIASES.get(key, header)


def rows_as_dicts(ws):
    headers = [canonical_header(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for excel_row in ws.iter_rows(min_row=2, values_only=True):
        if not any(value is not None and str(value).strip() for value in excel_row):
            continue
        row = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            row[header] = excel_row[index] if index < len(excel_row) else None
        rows.append(row)
    return rows


def main() -> int:
    if len(sys.argv) != 2:
        print("Usage: inspect_xlsform.py input.xlsx", file=sys.stderr)
        return 2

    input_path = Path(sys.argv[1])
    wb = load_workbook(input_path, data_only=False)
    if "survey" not in wb.sheetnames:
        raise SystemExit("Missing required XLSForm sheet: survey")

    survey_rows = rows_as_dicts(wb["survey"])
    settings_rows = rows_as_dicts(wb["settings"]) if "settings" in wb.sheetnames else []
    settings = settings_rows[0] if settings_rows else {}

    variables = []
    seen = set()
    for row in survey_rows:
        name = str(row.get("name") or "").strip()
        if not name or name in seen:
            continue
        seen.add(name)
        variables.append(
            {
                "name": name,
                "type": str(row.get("type") or "").strip(),
                "label": str(row.get("label") or "").strip(),
            }
        )

    payload = {
        "title": str(settings.get("form_title") or input_path.stem).strip(),
        "formId": str(settings.get("form_id") or input_path.stem).strip(),
        "variableCount": len(variables),
        "variables": variables,
    }
    print(json.dumps(payload, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
